import tempfile
import unittest
from pathlib import Path

import openpyxl
from docx import Document

from app.chunking import chunk_documents
from app.parsers import parse_docx, parse_xlsx


class TestTableChunking(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_xlsx_header_injection(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Assay"
        ws.append(["Crude Name", "API Gravity", "Sulfur wt%"])
        ws.append(["Arab Light", "33.2", "1.78"])
        ws.append(["Maya", "21.8", "3.40"])

        xlsx_path = Path(self.temp_dir.name) / "test_assay.xlsx"
        wb.save(str(xlsx_path))
        wb.close()

        records = parse_xlsx(xlsx_path)
        self.assertGreaterEqual(len(records), 3)

        # Row 2 (Arab Light) must have column headers injected
        row2 = records[1]
        self.assertTrue(row2.get("is_table"))
        self.assertIn("Crude Name: Arab Light", row2["text"])
        self.assertIn("API Gravity: 33.2", row2["text"])
        self.assertIn("Sulfur wt%: 1.78", row2["text"])

        # Chunker must preserve is_table flag
        chunks = chunk_documents(records)
        self.assertTrue(any(c.get("is_table") for c in chunks))

    def test_docx_table_header_injection(self):
        doc = Document()
        doc.add_heading("Equipment Limits", level=1)
        table = doc.add_table(rows=3, cols=3)
        headers = ["Equipment Tag", "Normal Pressure", "Trip Limit"]
        for i, h in enumerate(headers):
            table.cell(0, i).text = h

        row1 = ["BDV-701", "170 bar", "188 bar"]
        for i, val in enumerate(row1):
            table.cell(1, i).text = val

        row2 = ["PT-101", "3.0 bar", "2.2 bar"]
        for i, val in enumerate(row2):
            table.cell(2, i).text = val

        docx_path = Path(self.temp_dir.name) / "test_table.docx"
        doc.save(str(docx_path))

        records = parse_docx(docx_path)
        self.assertEqual(len(records), 1)
        text = records[0]["text"]
        self.assertIn("Equipment Tag: BDV-701", text)
        self.assertIn("Trip Limit: 188 bar", text)
        self.assertIn("Equipment Tag: PT-101", text)


if __name__ == "__main__":
    unittest.main()
