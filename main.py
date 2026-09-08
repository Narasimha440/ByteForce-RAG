from fastapi import FastAPI, Depends, HTTPException, status
from fastapi.security import OAuth2PasswordRequestForm
from sqlalchemy.orm import Session
from pydantic import BaseModel, Field, field_validator, model_validator
from docx import Document as DocxDocument
import io
import re
import logging
from typing import List, Optional, Any
from fastapi.responses import StreamingResponse

from database import Base, engine, get_db
import models
from auth import hash_password, verify_password, create_access_token, get_current_user
from api_clients import call_embed, call_vision, call_coder
from app.agent.agent import Agent
from generators.pptx_generator import generate_pptx
from generators.xlsx_generator import generate_xlsx

Base.metadata.create_all(bind=engine)
app = FastAPI(title="SIH26117 Backend")
agent = Agent()

class RegisterRequest(BaseModel):
    username: str
    password: str

@app.post("/api/auth/register")
def register(req: RegisterRequest, db: Session = Depends(get_db)):
    existing = db.query(models.User).filter(models.User.username == req.username).first()
    if existing:
        raise HTTPException(400, "Username already exists")
    user = models.User(username=req.username, hashed_password=hash_password(req.password))
    db.add(user)
    db.commit()
    return {"message": "User created"}

@app.post("/api/auth/login")
def login(form: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    user = db.query(models.User).filter(models.User.username == form.username).first()
    if not user or not verify_password(form.password, user.hashed_password):
        raise HTTPException(401, "Invalid credentials")
    token = create_access_token({"sub": user.username})
    return {"access_token": token, "token_type": "bearer"}


class ChatRequest(BaseModel):
    message: str
class AgentQueryRequest(BaseModel):
    question: str
@app.post("/api/agent/query")
def agent_query(
    req: AgentQueryRequest,
    current_user: models.User = Depends(get_current_user)
):
    try:
        result = ""

        for event in agent.run_stream(req.question):
            if event.get("type") == "token":
                result += event.get("content", "")

            elif event.get("type") == "done":
                result = event.get("result", {}).get(
                    "answer",
                    result
                )

        return {
            "answer": result
        }

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Agent execution failed: {str(e)}"
        )
@app.post("/api/chat")
def chat(
    req: ChatRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    steps = []
    final_answer = ""

    try:
        for event in agent.run_stream(req.message):

            event_type = event.get("type")

            if event_type == "step":
                steps.append(event.get("message", ""))

            elif event_type == "classification":
                steps.append(
                    f"Intent classified as {event.get('intent', 'unknown')}"
                )

            elif event_type == "plan":
                steps.append("Execution plan created")

            elif event_type == "replan":
                steps.append(
                    f"Dynamic replan {event.get('attempt', 0)}"
                )

            elif event_type == "token":
                final_answer += event.get("content", "")

            elif event_type == "done":
                result = event.get("result", {})
                final_answer = result.get(
                    "answer",
                    final_answer
                )

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Agent execution failed: {str(e)}"
        )

    session = models.ChatSession(
        user_id=current_user.id,
        message=req.message,
        final_answer=final_answer
    )

    db.add(session)
    db.commit()

    return {
        "final_answer": final_answer,
        "steps": steps
    }


@app.get("/api/history")
def history(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    sessions = db.query(models.ChatSession).filter(models.ChatSession.user_id == current_user.id).all()
    return [{"message": s.message, "answer": s.final_answer, "created_at": s.created_at} for s in sessions]


class DocxRequest(BaseModel):
    title: str
    content: str

@app.post("/api/generate/docx")
def generate_docx(req: DocxRequest, current_user: models.User = Depends(get_current_user)):
    doc = DocxDocument()
    doc.add_heading(req.title, level=1)
    doc.add_paragraph(req.content)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                             headers={"Content-Disposition": f"attachment; filename={req.title}.docx"})


def _get_safe_filename(name: Optional[str], default_name: str, ext: str) -> str:
    if not name or not name.strip():
        return f"{default_name}.{ext}"
    cleaned = re.sub(r'[\\/*?:"<>|]', "", name.strip()).strip().replace(" ", "_")
    return f"{cleaned}.{ext}" if cleaned else f"{default_name}.{ext}"


class SlideItem(BaseModel):
    title: str = Field(..., description="Slide title")
    content: Optional[str] = Field(None, description="Narrative content for the slide")
    bullet_points: Optional[List[str]] = Field(None, description="Bullet points for the slide")

    @field_validator("title")
    @classmethod
    def validate_title(cls, v: str) -> str:
        if not v or not v.strip():
            raise ValueError("Slide title cannot be empty.")
        return v.strip()


class PPTXRequest(BaseModel):
    title: str = Field(..., description="Presentation title")
    subtitle: Optional[str] = Field(None, description="Optional presentation subtitle")
    slides: List[SlideItem] = Field(..., description="List of presentation slides")

    @field_validator("title")
    @classmethod
    def validate_title(cls, v: str) -> str:
        if not v or not v.strip():
            raise ValueError("Presentation title cannot be empty.")
        return v.strip()

    @field_validator("slides")
    @classmethod
    def validate_slides(cls, v: List[SlideItem]) -> List[SlideItem]:
        if not v:
            raise ValueError("Slides list cannot be empty.")
        return v


class XLSXRequest(BaseModel):
    title: Optional[str] = Field(None, description="Optional title of the spreadsheet")
    headers: List[str] = Field(..., description="List of column headers")
    rows: List[List[Any]] = Field(..., description="Tabular rows of data")

    @model_validator(mode="after")
    def validate_tabular_structure(self):
        if not self.headers:
            raise ValueError("Headers list cannot be empty.")
        for h in self.headers:
            if not str(h).strip():
                raise ValueError("Header titles cannot be empty.")
        for idx, row in enumerate(self.rows):
            if len(row) != len(self.headers):
                raise ValueError(
                    f"Row {idx + 1} has {len(row)} items; expected {len(self.headers)} to match headers."
                )
        return self


@app.post("/api/generate/pptx")
def generate_pptx_endpoint(
    req: PPTXRequest,
    current_user: models.User = Depends(get_current_user)
):
    try:
        slides_data = [
            {
                "title": s.title,
                "content": s.content,
                "bullet_points": s.bullet_points
            }
            for s in req.slides
        ]
        buffer = generate_pptx(
            title=req.title,
            slides=slides_data,
            subtitle=req.subtitle
        )
        filename = _get_safe_filename(req.title, "presentation", "pptx")
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve))
    except Exception as e:
        logging.getLogger("uvicorn.error").exception("Failed to generate PPTX presentation")
        raise HTTPException(
            status_code=500,
            detail="Failed to generate PowerPoint presentation."
        )


@app.post("/api/generate/xlsx")
def generate_xlsx_endpoint(
    req: XLSXRequest,
    current_user: models.User = Depends(get_current_user)
):
    try:
        buffer = generate_xlsx(
            headers=req.headers,
            rows=req.rows,
            title=req.title
        )
        filename = _get_safe_filename(req.title, "report", "xlsx")
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve))
    except Exception as e:
        logging.getLogger("uvicorn.error").exception("Failed to generate XLSX spreadsheet")
        raise HTTPException(
            status_code=500,
            detail="Failed to generate Excel spreadsheet."
        )

@app.get("/")
def root():
    return {"status": "SIH26117 backend running"}