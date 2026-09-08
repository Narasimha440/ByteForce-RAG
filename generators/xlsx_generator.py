"""
Excel (.xlsx) spreadsheet generator.
Creates clean, professional, auto-formatted worksheets in memory using openpyxl.
"""

import io
import re
import logging
from typing import List, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


def generate_xlsx(
    headers: List[str],
    rows: List[List[Any]],
    title: Optional[str] = None
) -> io.BytesIO:
    """
    Generate an Excel (.xlsx) file in memory from structured tabular data.

    :param headers: List of column header names.
    :param rows: List of row records; each record must match the length of headers.
    :param title: Optional title printed at the top of the sheet.
    :return: io.BytesIO buffer positioned at 0 containing the valid XLSX file.
    """
    if not headers:
        raise ValueError("Headers list cannot be empty.")

    # Validate each row length against headers
    for idx, row in enumerate(rows):
        if len(row) != len(headers):
            raise ValueError(
                f"Row {idx + 1} has {len(row)} values; expected {len(headers)} to match headers."
            )

    wb = openpyxl.Workbook()
    ws = wb.active

    # Configure Sheet Tab Name
    if title and title.strip():
        # Excel sheet titles cannot exceed 31 chars or contain : \ / ? * [ ]
        clean_tab_name = re.sub(r"[:\\/?*\[\]]", "_", title.strip())[:31]
        ws.title = clean_tab_name or "Report"
    else:
        ws.title = "Report"

    # Styling Elements
    title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    data_font = Font(name="Calibri", size=11, color="111827")

    thin_border_side = Side(style="thin", color="CBD5E1")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )

    current_row = 1

    # 1. Title Row (Optional)
    if title and title.strip():
        ws.cell(row=current_row, column=1, value=title.strip())
        ws.cell(row=current_row, column=1).font = title_font
        ws.row_dimensions[current_row].height = 28
        current_row += 1

        # Spacer row
        ws.row_dimensions[current_row].height = 10
        current_row += 1

    # 2. Headers Row
    header_row_idx = current_row
    ws.row_dimensions[header_row_idx].height = 24

    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=str(header_text))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border

    current_row += 1

    # 3. Data Rows
    for row_data in rows:
        ws.row_dimensions[current_row].height = 20
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = data_font
            cell.alignment = Alignment(vertical="center")
            cell.border = cell_border
        current_row += 1

    # 4. Auto-adjust Column Widths (ignoring title row to prevent ballooning)
    for col in ws.iter_cols(min_col=1, max_col=len(headers)):
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row < header_row_idx:
                continue
            val_str = str(cell.value or "")
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # 5. Freeze Panes directly below header row
    ws.freeze_panes = f"A{header_row_idx + 1}"

    # 6. Save into memory buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
