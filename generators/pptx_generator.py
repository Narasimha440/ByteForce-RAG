"""
PowerPoint (.pptx) presentation generator.
Creates clean, professional slides in memory using python-pptx.
"""

import io
import logging
from typing import List, Dict, Any, Optional, Union
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

logger = logging.getLogger(__name__)

# Styling Constants
COLOR_TITLE = RGBColor(0x11, 0x18, 0x27)       # Deep slate gray
COLOR_SUBTITLE = RGBColor(0x4B, 0x55, 0x63)    # Muted gray
COLOR_HEADING = RGBColor(0x1E, 0x3A, 0x8A)     # Professional navy blue
COLOR_BODY = RGBColor(0x37, 0x41, 0x51)        # Charcoal body text

FONT_TITLE_SIZE = Pt(38)
FONT_SUBTITLE_SIZE = Pt(20)
FONT_SLIDE_TITLE_SIZE = Pt(28)
FONT_BODY_SIZE = Pt(18)
FONT_BULLET_SIZE = Pt(16)


def generate_pptx(
    title: str,
    slides: List[Dict[str, Any]],
    subtitle: Optional[str] = None
) -> io.BytesIO:
    """
    Generate a PowerPoint presentation in memory based on structured input.

    :param title: Main title for the presentation.
    :param slides: List of slide dicts, each with:
                   - 'title' (str): Title for the slide.
                   - 'content' (Optional[str]): Narrative body text.
                   - 'bullet_points' (Optional[List[str]]): List of bullet items.
    :param subtitle: Optional subtitle for the title slide.
    :return: io.BytesIO buffer positioned at 0 containing the valid PPTX file.
    """
    if not title or not title.strip():
        raise ValueError("Presentation title cannot be empty.")

    prs = Presentation()

    # -------------------------------------------------------------------------
    # 1. TITLE SLIDE (Layout 0)
    # -------------------------------------------------------------------------
    title_slide_layout = prs.slide_layouts[0]
    title_slide = prs.slides.add_slide(title_slide_layout)

    # Set Title
    if title_slide.shapes.title:
        title_slide.shapes.title.text = title.strip()
        for p in title_slide.shapes.title.text_frame.paragraphs:
            p.font.size = FONT_TITLE_SIZE
            p.font.bold = True
            p.font.color.rgb = COLOR_TITLE

    # Set Subtitle (if provided)
    if subtitle and subtitle.strip() and len(title_slide.placeholders) > 1:
        subtitle_shape = title_slide.placeholders[1]
        subtitle_shape.text = subtitle.strip()
        for p in subtitle_shape.text_frame.paragraphs:
            p.font.size = FONT_SUBTITLE_SIZE
            p.font.color.rgb = COLOR_SUBTITLE

    # -------------------------------------------------------------------------
    # 2. CONTENT SLIDES (Layout 1: Title and Content)
    # -------------------------------------------------------------------------
    content_layout = prs.slide_layouts[1]

    for slide_data in slides:
        slide_title = str(slide_data.get("title") or "Slide").strip()
        slide_content = slide_data.get("content")
        bullet_points = slide_data.get("bullet_points")

        slide = prs.slides.add_slide(content_layout)

        # Slide title
        if slide.shapes.title:
            slide.shapes.title.text = slide_title
            for p in slide.shapes.title.text_frame.paragraphs:
                p.font.size = FONT_SLIDE_TITLE_SIZE
                p.font.bold = True
                p.font.color.rgb = COLOR_HEADING

        # Body / Content placeholder
        if len(slide.placeholders) > 1:
            body_shape = slide.placeholders[1]
            tf = body_shape.text_frame
            tf.word_wrap = True

            clean_content = slide_content.strip() if slide_content and isinstance(slide_content, str) else ""
            clean_bullets = [
                str(b).strip() for b in bullet_points if b is not None and str(b).strip()
            ] if bullet_points and isinstance(bullet_points, list) else []

            if clean_content and clean_bullets:
                # Content followed by bullet points
                tf.text = clean_content
                p0 = tf.paragraphs[0]
                p0.font.size = FONT_BODY_SIZE
                p0.font.color.rgb = COLOR_BODY
                p0.level = 0

                for bp in clean_bullets:
                    p = tf.add_paragraph()
                    p.text = bp
                    p.font.size = FONT_BULLET_SIZE
                    p.font.color.rgb = COLOR_BODY
                    p.level = 1

            elif clean_content and not clean_bullets:
                # Content only
                tf.text = clean_content
                p0 = tf.paragraphs[0]
                p0.font.size = FONT_BODY_SIZE
                p0.font.color.rgb = COLOR_BODY
                p0.level = 0

            elif not clean_content and clean_bullets:
                # Bullets only
                tf.text = clean_bullets[0]
                p0 = tf.paragraphs[0]
                p0.font.size = FONT_BODY_SIZE
                p0.font.color.rgb = COLOR_BODY
                p0.level = 0

                for bp in clean_bullets[1:]:
                    p = tf.add_paragraph()
                    p.text = bp
                    p.font.size = FONT_BODY_SIZE
                    p.font.color.rgb = COLOR_BODY
                    p.level = 0

            else:
                # Neither provided
                tf.text = ""

    # -------------------------------------------------------------------------
    # 3. SAVE TO IN-MEMORY BUFFER
    # -------------------------------------------------------------------------
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer
